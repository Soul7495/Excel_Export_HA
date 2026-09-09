"""Device Entity XLSX Export integration."""

from __future__ import annotations

from datetime import datetime
import json
import logging
from pathlib import Path
import re
from typing import Any

import voluptuous as vol
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from homeassistant.components import persistent_notification
from homeassistant.config_entries import ConfigEntry
from homeassistant.core import HomeAssistant, ServiceCall, ServiceResponse, SupportsResponse
from homeassistant.exceptions import ServiceValidationError
from homeassistant.helpers import config_validation as cv
from homeassistant.helpers import device_registry as dr
from homeassistant.helpers import entity_registry as er
from homeassistant.helpers.typing import ConfigType

from .const import (
    ATTR_DEVICE_ID,
    ATTR_FILENAME,
    ATTR_INCLUDE_ATTRIBUTES,
    ATTR_INCLUDE_DISABLED,
    DEFAULT_INCLUDE_ATTRIBUTES,
    DEFAULT_INCLUDE_DISABLED,
    DOMAIN,
    EXPORT_DIR,
    SERVICE_EXPORT_DEVICE,
)

_LOGGER = logging.getLogger(__name__)

SERVICE_SCHEMA = vol.Schema(
    {
        vol.Required(ATTR_DEVICE_ID): cv.string,
        vol.Optional(ATTR_FILENAME): cv.string,
        vol.Optional(
            ATTR_INCLUDE_DISABLED,
            default=DEFAULT_INCLUDE_DISABLED,
        ): cv.boolean,
        vol.Optional(
            ATTR_INCLUDE_ATTRIBUTES,
            default=DEFAULT_INCLUDE_ATTRIBUTES,
        ): cv.boolean,
    }
)


def _enum_value(value: Any) -> str | None:
    """Convert enum-like values to strings."""
    if value is None:
        return None
    return str(getattr(value, "value", value))


def _excel_safe(value: Any) -> Any:
    """Make a value safe and compact enough for an Excel cell."""
    if value is None:
        return ""
    if isinstance(value, (bool, int, float, datetime)):
        return value

    text = str(value)
    if len(text) > 32700:
        text = text[:32680] + " … [gekürzt]"

    # Prevent accidental formula execution for user/integration-provided strings.
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _slugify(value: str) -> str:
    """Create a filesystem-safe filename fragment."""
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9äöüß_-]+", "_", value, flags=re.IGNORECASE)
    value = re.sub(r"_+", "_", value).strip("_.-")
    return value or "device"


def _normalise_filename(filename: str | None, device_name: str) -> str:
    """Return a safe XLSX filename."""
    if filename:
        candidate = Path(filename).name.strip()
        if not candidate.lower().endswith(".xlsx"):
            candidate += ".xlsx"
        stem = _slugify(Path(candidate).stem)
        return f"{stem}.xlsx"

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{_slugify(device_name)}_{timestamp}.xlsx"


def _unique_path(directory: Path, filename: str) -> Path:
    """Avoid overwriting an existing export."""
    path = directory / filename
    if not path.exists():
        return path

    stem = path.stem
    suffix = path.suffix
    counter = 2
    while True:
        candidate = directory / f"{stem}_{counter}{suffix}"
        if not candidate.exists():
            return candidate
        counter += 1


def _json_attributes(attributes: dict[str, Any]) -> str:
    """Serialize attributes without failing on unusual values."""
    return json.dumps(
        attributes,
        ensure_ascii=False,
        sort_keys=True,
        default=str,
    )


def _write_workbook(
    path: Path,
    device_data: dict[str, Any],
    rows: list[dict[str, Any]],
    include_attributes: bool,
) -> None:
    """Write workbook synchronously in an executor thread."""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Entitäten"

    columns: list[tuple[str, str]] = [
        ("entity_id", "Entity ID"),
        ("name", "Name"),
        ("state", "Zustand"),
        ("unit", "Einheit"),
        ("domain", "Domain"),
        ("platform", "Integration / Plattform"),
        ("device_class", "Device Class"),
        ("state_class", "State Class"),
        ("entity_category", "Entity Category"),
        ("disabled_by", "Deaktiviert durch"),
        ("hidden_by", "Ausgeblendet durch"),
        ("available", "Aktuell geladen"),
        ("last_changed", "Letzte Zustandsänderung"),
        ("last_updated", "Letzte Aktualisierung"),
        ("unique_id", "Unique ID"),
        ("config_entry_id", "Config Entry ID"),
        ("area_id", "Area ID"),
    ]
    if include_attributes:
        columns.append(("attributes", "Attribute (JSON)"))

    ws.append([title for _, title in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top")

    for row in rows:
        ws.append([_excel_safe(row.get(key)) for key, _ in columns])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "Entity ID": 42,
        "Name": 34,
        "Zustand": 22,
        "Einheit": 16,
        "Domain": 16,
        "Integration / Plattform": 24,
        "Device Class": 22,
        "State Class": 20,
        "Entity Category": 20,
        "Deaktiviert durch": 20,
        "Ausgeblendet durch": 20,
        "Aktuell geladen": 16,
        "Letzte Zustandsänderung": 25,
        "Letzte Aktualisierung": 25,
        "Unique ID": 42,
        "Config Entry ID": 38,
        "Area ID": 26,
        "Attribute (JSON)": 70,
    }
    for index, (_, title) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(title, 22)

    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    device_ws = workbook.create_sheet("Gerät")
    device_ws.append(["Feld", "Wert"])
    device_ws["A1"].font = Font(bold=True)
    device_ws["B1"].font = Font(bold=True)
    for key, value in device_data.items():
        device_ws.append([key, _excel_safe(value)])
    device_ws.column_dimensions["A"].width = 26
    device_ws.column_dimensions["B"].width = 70
    device_ws.freeze_panes = "A2"

    info_ws = workbook.create_sheet("Exportinfo")
    info_ws.append(["Feld", "Wert"])
    info_ws["A1"].font = Font(bold=True)
    info_ws["B1"].font = Font(bold=True)
    info_ws.append(["Exportiert am", datetime.now().astimezone().isoformat(timespec="seconds")])
    info_ws.append(["Anzahl Entitäten", len(rows)])
    info_ws.append(["Deaktivierte Entitäten enthalten", any(row["disabled_by"] for row in rows)])
    info_ws.append(["Attribute enthalten", include_attributes])
    info_ws.column_dimensions["A"].width = 34
    info_ws.column_dimensions["B"].width = 70

    workbook.save(path)


async def async_setup(hass: HomeAssistant, config: ConfigType) -> bool:
    """Set up Device Entity XLSX Export."""

    if hass.services.has_service(DOMAIN, SERVICE_EXPORT_DEVICE):
        return True

    async def async_handle_export(call: ServiceCall) -> ServiceResponse | None:
        """Export all entities assigned to one Home Assistant device."""
        device_id: str = call.data[ATTR_DEVICE_ID]
        include_disabled: bool = call.data[ATTR_INCLUDE_DISABLED]
        include_attributes: bool = call.data[ATTR_INCLUDE_ATTRIBUTES]
        requested_filename: str | None = call.data.get(ATTR_FILENAME)

        device_registry = dr.async_get(hass)
        device = device_registry.async_get(device_id)
        if device is None:
            raise ServiceValidationError(
                translation_domain=DOMAIN,
                translation_key="device_not_found",
                translation_placeholders={"device_id": device_id},
            )

        entity_registry = er.async_get(hass)
        registry_entries = er.async_entries_for_device(
            entity_registry, device_id, include_disabled_entities=True
        )

        device_name = device.name_by_user or device.name or device.id
        rows: list[dict[str, Any]] = []

        for entry in sorted(registry_entries, key=lambda item: item.entity_id):
            if not include_disabled and entry.disabled_by is not None:
                continue

            state = hass.states.get(entry.entity_id)
            attributes = dict(state.attributes) if state is not None else {}
            name = (
                entry.name
                or entry.original_name
                or attributes.get("friendly_name")
                or entry.entity_id
            )

            rows.append(
                {
                    "entity_id": entry.entity_id,
                    "name": name,
                    "state": state.state if state is not None else "",
                    "unit": attributes.get("unit_of_measurement")
                    or entry.unit_of_measurement,
                    "domain": entry.domain,
                    "platform": entry.platform,
                    "device_class": attributes.get("device_class")
                    or entry.device_class
                    or entry.original_device_class,
                    "state_class": attributes.get("state_class"),
                    "entity_category": _enum_value(entry.entity_category),
                    "disabled_by": _enum_value(entry.disabled_by),
                    "hidden_by": _enum_value(entry.hidden_by),
                    "available": state is not None,
                    "last_changed": state.last_changed.isoformat(timespec="seconds")
                    if state is not None
                    else "",
                    "last_updated": state.last_updated.isoformat(timespec="seconds")
                    if state is not None
                    else "",
                    "unique_id": entry.unique_id,
                    "config_entry_id": entry.config_entry_id,
                    "area_id": entry.area_id,
                    "attributes": _json_attributes(attributes)
                    if include_attributes
                    else "",
                }
            )

        if not rows:
            raise ServiceValidationError(
                translation_domain=DOMAIN,
                translation_key="no_entities",
                translation_placeholders={"device_name": device_name},
            )

        export_directory = Path(hass.config.path(EXPORT_DIR))
        await hass.async_add_executor_job(export_directory.mkdir, 0o755, True, True)

        filename = _normalise_filename(requested_filename, device_name)
        output_path = await hass.async_add_executor_job(
            _unique_path,
            export_directory,
            filename,
        )

        device_data = {
            "Name": device_name,
            "Home Assistant Device ID": device.id,
            "Hersteller": device.manufacturer,
            "Modell": device.model,
            "Modell-ID": getattr(device, "model_id", None),
            "Seriennummer": device.serial_number,
            "Software-Version": device.sw_version,
            "Hardware-Version": device.hw_version,
            "Area ID": device.area_id,
            "Anzahl exportierter Entitäten": len(rows),
        }

        await hass.async_add_executor_job(
            _write_workbook,
            output_path,
            device_data,
            rows,
            include_attributes,
        )

        relative_path = f"/{EXPORT_DIR}/{output_path.name}"
        persistent_notification.async_create(
            hass,
            (
                f"**{device_name}** wurde exportiert.\n\n"
                f"Datei: `{output_path}`\n\n"
                "Die Datei liegt im Home-Assistant-Konfigurationsordner unter "
                f"`{EXPORT_DIR}/`."
            ),
            title="Excel-Export abgeschlossen",
            notification_id=f"{DOMAIN}_last_export",
        )

        _LOGGER.info(
            "Exported %s entities for device %s to %s",
            len(rows),
            device_name,
            output_path,
        )

        response: ServiceResponse = {
            "device": device_name,
            "device_id": device.id,
            "entity_count": len(rows),
            "filename": output_path.name,
            "path": str(output_path),
            "config_relative_path": relative_path,
        }
        if call.return_response:
            return response
        return None

    hass.services.async_register(
        DOMAIN,
        SERVICE_EXPORT_DEVICE,
        async_handle_export,
        schema=SERVICE_SCHEMA,
        supports_response=SupportsResponse.OPTIONAL,
    )

    return True


async def async_setup_entry(hass: HomeAssistant, entry: ConfigEntry) -> bool:
    """Set up from a config entry."""
    return True


async def async_unload_entry(hass: HomeAssistant, entry: ConfigEntry) -> bool:
    """Unload a config entry."""
    return True
