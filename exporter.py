"""Collect Home Assistant registry data and create XLSX workbooks."""
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import re
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from homeassistant.core import HomeAssistant
from homeassistant.exceptions import ServiceValidationError
from homeassistant.helpers import device_registry as dr, entity_registry as er
from .const import DOMAIN, EXPORT_DIR

def _enum(value):
    return None if value is None else str(getattr(value, "value", value))

def _safe(value):
    if value is None: return ""
    if isinstance(value, (bool, int, float, datetime)): return value
    text = str(value)[:32700]
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text

def normalise_filename(filename, device_name):
    value = Path(filename).stem if filename else f"{device_name}_{datetime.now():%Y%m%d_%H%M%S}"
    value = re.sub(r"[^a-z0-9äöüß_-]+", "_", value.strip().lower(), flags=re.I)
    return f"{re.sub(r'_+', '_', value).strip('_.-') or 'device'}.xlsx"

def list_devices(hass):
    devices, entities = dr.async_get(hass), er.async_get(hass)
    counts = {}
    for entry in entities.entities.values():
        if entry.device_id: counts[entry.device_id] = counts.get(entry.device_id, 0) + 1
    result = []
    for device in devices.devices.values():
        if not counts.get(device.id): continue
        result.append({"id": device.id, "name": device.name_by_user or device.name or device.id, "manufacturer": device.manufacturer or "", "model": device.model or "", "area_id": device.area_id, "entity_count": counts[device.id]})
    return sorted(result, key=lambda item: item["name"].casefold())

def collect_device(hass, device_id, include_disabled=True):
    device = dr.async_get(hass).async_get(device_id)
    if device is None:
        raise ServiceValidationError(translation_domain=DOMAIN, translation_key="device_not_found", translation_placeholders={"device_id": device_id})
    rows = []
    entries = er.async_entries_for_device(er.async_get(hass), device_id, include_disabled_entities=True)
    for entry in sorted(entries, key=lambda item: item.entity_id):
        if not include_disabled and entry.disabled_by is not None: continue
        state = hass.states.get(entry.entity_id)
        attrs = dict(state.attributes) if state else {}
        rows.append({"entity_id": entry.entity_id, "name": entry.name or entry.original_name or attrs.get("friendly_name") or entry.entity_id, "state": state.state if state else "", "unit": attrs.get("unit_of_measurement") or entry.unit_of_measurement, "domain": entry.domain, "platform": entry.platform, "device_class": attrs.get("device_class") or entry.device_class or entry.original_device_class, "state_class": attrs.get("state_class"), "entity_category": _enum(entry.entity_category), "disabled_by": _enum(entry.disabled_by), "hidden_by": _enum(entry.hidden_by), "available": state is not None, "last_changed": state.last_changed.isoformat(timespec="seconds") if state else "", "last_updated": state.last_updated.isoformat(timespec="seconds") if state else "", "unique_id": entry.unique_id, "config_entry_id": entry.config_entry_id, "area_id": entry.area_id, "attributes": json.dumps(attrs, ensure_ascii=False, sort_keys=True, default=str)})
    name = device.name_by_user or device.name or device.id
    if not rows:
        raise ServiceValidationError(translation_domain=DOMAIN, translation_key="no_entities", translation_placeholders={"device_name": name})
    metadata = {"Name": name, "Home Assistant Device ID": device.id, "Hersteller": device.manufacturer, "Modell": device.model, "Modell-ID": getattr(device, "model_id", None), "Seriennummer": device.serial_number, "Software-Version": device.sw_version, "Hardware-Version": device.hw_version, "Area ID": device.area_id, "Anzahl exportierter Entitäten": len(rows)}
    return metadata, rows

def build_workbook(metadata, rows, include_attributes):
    book = Workbook(); sheet = book.active; sheet.title = "Entitäten"
    columns = [("entity_id","Entity ID"),("name","Name"),("state","Zustand"),("unit","Einheit"),("domain","Domain"),("platform","Integration / Plattform"),("device_class","Device Class"),("state_class","State Class"),("entity_category","Entity Category"),("disabled_by","Deaktiviert durch"),("hidden_by","Ausgeblendet durch"),("available","Aktuell geladen"),("last_changed","Letzte Zustandsänderung"),("last_updated","Letzte Aktualisierung"),("unique_id","Unique ID"),("config_entry_id","Config Entry ID"),("area_id","Area ID")]
    if include_attributes: columns.append(("attributes", "Attribute (JSON)"))
    sheet.append([title for _, title in columns])
    for cell in sheet[1]: cell.font = Font(bold=True)
    for row in rows: sheet.append([_safe(row.get(key)) for key, _ in columns])
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    widths = [42,34,22,16,16,24,22,20,20,20,20,16,25,25,42,38,26,70]
    for i in range(1, len(columns)+1): sheet.column_dimensions[get_column_letter(i)].width = widths[i-1]
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells: cell.alignment = Alignment(vertical="top")
    device_sheet = book.create_sheet("Gerät"); device_sheet.append(["Feld", "Wert"])
    for cell in device_sheet[1]: cell.font = Font(bold=True)
    for key, value in metadata.items(): device_sheet.append([key, _safe(value)])
    device_sheet.column_dimensions["A"].width = 34; device_sheet.column_dimensions["B"].width = 70
    info = book.create_sheet("Exportinfo"); info.append(["Feld", "Wert"]); info.append(["Exportiert am", datetime.now().astimezone().isoformat(timespec="seconds")]); info.append(["Anzahl Entitäten", len(rows)]); info.append(["Attribute enthalten", include_attributes])
    for cell in info[1]: cell.font = Font(bold=True)
    output = BytesIO(); book.save(output); return output.getvalue()

async def async_build_export(hass: HomeAssistant, device_id: str, filename: str | None, include_disabled: bool, include_attributes: bool):
    metadata, rows = collect_device(hass, device_id, include_disabled)
    content = await hass.async_add_executor_job(build_workbook, metadata, rows, include_attributes)
    return normalise_filename(filename, metadata["Name"]), content, len(rows), metadata["Name"]

async def async_export_to_disk(hass: HomeAssistant, **options: Any):
    filename, content, count, name = await async_build_export(hass, **options)
    directory = Path(hass.config.path(EXPORT_DIR)); await hass.async_add_executor_job(directory.mkdir, 0o755, True, True)
    path = directory / filename; counter = 2
    while path.exists(): path = directory / f"{Path(filename).stem}_{counter}.xlsx"; counter += 1
    await hass.async_add_executor_job(path.write_bytes, content)
    return {"device": name, "device_id": options["device_id"], "entity_count": count, "filename": path.name, "path": str(path)}
